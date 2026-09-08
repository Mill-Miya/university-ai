from __future__ import annotations

import json
from datetime import UTC, datetime
from pathlib import Path

import pytest
from docx import Document as DocxDocument

from university_ai.database.database import Database
from university_ai.database.migrations import migrate
from university_ai.database.models import Document, ExtractionStatus
from university_ai.database.repository import DocumentRepository
from university_ai.documents.import_service import DocumentImportError, DocumentImportService
from university_ai.ui.documents import DocumentPresenter


@pytest.fixture
def document_services(tmp_path):
    database = Database(tmp_path / "app.sqlite3")
    connection = database.connect()
    migrate(connection)
    repository = DocumentRepository(connection)
    service = DocumentImportService(repository, tmp_path / "managed-documents")
    yield repository, service
    database.stop()


def _write_text_pdf(path: Path, text: str) -> None:
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 5 0 R >> >> /Contents 4 0 R >>",
        f"<< /Length {len(f'BT /F1 12 Tf 72 720 Td ({text}) Tj ET'.encode('ascii'))} >>\nstream\nBT /F1 12 Tf 72 720 Td ({text}) Tj ET\nendstream".encode("ascii"),
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    payload = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for index, value in enumerate(objects, 1):
        offsets.append(len(payload))
        payload.extend(f"{index} 0 obj\n".encode())
        payload.extend(value)
        payload.extend(b"\nendobj\n")
    startxref = len(payload)
    payload.extend(f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode())
    payload.extend(b"".join(f"{offset:010d} 00000 n \n".encode() for offset in offsets[1:]))
    payload.extend(f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{startxref}\n%%EOF\n".encode())
    path.write_bytes(payload)


def test_txt_import_copies_source_extracts_text_and_persists(document_services, tmp_path):
    repository, service = document_services
    source = tmp_path / "notes.txt"
    source.write_text("Local university notes", encoding="utf-8")
    original = source.read_bytes()

    imported = service.import_file(source).document

    assert imported.file_type == "TXT"
    assert imported.extraction_status is ExtractionStatus.EXTRACTED
    assert imported.extracted_text == "Local university notes"
    assert Path(imported.stored_path).read_bytes() == original
    assert source.read_bytes() == original
    assert repository.get(imported.id or 0) == imported


def test_document_is_available_after_database_restart(tmp_path):
    database = Database(tmp_path / "app.sqlite3")
    connection = database.connect(); migrate(connection)
    repository = DocumentRepository(connection)
    service = DocumentImportService(repository, tmp_path / "managed-documents")
    source = tmp_path / "restart.txt"; source.write_text("survives", encoding="utf-8")
    imported = service.import_file(source).document
    database.stop()

    reopened = Database(tmp_path / "app.sqlite3")
    reopened_connection = reopened.connect(); migrate(reopened_connection)
    assert DocumentRepository(reopened_connection).get(imported.id or 0) == imported
    reopened.stop()


def test_pdf_and_docx_text_extraction(document_services, tmp_path):
    _repository, service = document_services
    pdf = tmp_path / "lecture.pdf"
    _write_text_pdf(pdf, "PDF lecture text")
    docx = tmp_path / "report.docx"
    writer = DocxDocument(); writer.add_paragraph("DOCX lecture text"); writer.save(docx)

    pdf_result = service.import_file(pdf).document
    docx_result = service.import_file(docx).document

    assert pdf_result.extraction_status is ExtractionStatus.EXTRACTED
    assert "PDF lecture text" in (pdf_result.extracted_text or "")
    assert docx_result.extraction_status is ExtractionStatus.EXTRACTED
    assert docx_result.extracted_text == "DOCX lecture text"


def test_same_content_is_deduplicated_by_sha256_not_filename(document_services, tmp_path):
    repository, service = document_services
    first = tmp_path / "week1.txt"; second = tmp_path / "renamed.txt"
    first.write_text("same bytes", encoding="utf-8"); second.write_text("same bytes", encoding="utf-8")

    initial = service.import_file(first)
    repeated = service.import_file(second)

    assert not initial.duplicate and repeated.duplicate
    assert repeated.document.id == initial.document.id
    assert len(repository.list()) == 1


def test_storage_path_uses_hash_to_avoid_filename_collisions(document_services, tmp_path):
    _repository, service = document_services
    first = tmp_path / "first" / "lecture.txt"; second = tmp_path / "second" / "lecture.txt"
    first.parent.mkdir(); second.parent.mkdir()
    first.write_text("first", encoding="utf-8"); second.write_text("second", encoding="utf-8")

    one = service.import_file(first).document
    two = service.import_file(second).document

    assert one.stored_path != two.stored_path
    assert Path(one.stored_path).read_text(encoding="utf-8") == "first"
    assert Path(two.stored_path).read_text(encoding="utf-8") == "second"


def test_unsupported_and_corrupted_files_do_not_create_successful_document(document_services, tmp_path):
    repository, service = document_services
    unsupported = tmp_path / "archive.zip"; unsupported.write_bytes(b"not used")
    broken = tmp_path / "broken.pdf"; broken.write_bytes(b"not a PDF")
    broken_docx = tmp_path / "broken.docx"; broken_docx.write_bytes(b"not a DOCX")

    with pytest.raises(DocumentImportError, match="unsupported"):
        service.import_file(unsupported)
    imported = service.import_file(broken).document
    imported_docx = service.import_file(broken_docx).document

    assert imported.extraction_status is ExtractionStatus.FAILED
    assert imported.extraction_error
    assert repository.get(imported.id or 0).extraction_status is ExtractionStatus.FAILED
    assert imported_docx.extraction_status is ExtractionStatus.FAILED


def test_encoding_failure_is_recorded_and_document_record_can_be_deleted(document_services, tmp_path):
    repository, service = document_services
    invalid_utf8 = tmp_path / "legacy.txt"; invalid_utf8.write_bytes(b"\xff\xfe")

    imported = service.import_file(invalid_utf8).document

    assert imported.extraction_status is ExtractionStatus.FAILED
    assert "utf" in (imported.extraction_error or "").lower()
    assert repository.delete(imported.id or 0)
    assert repository.get(imported.id or 0) is None


def test_image_is_saved_without_ocr_and_invalid_document_datetime_is_rejected(document_services, tmp_path):
    repository, service = document_services
    image = tmp_path / "slide.png"; image.write_bytes(b"png bytes")
    imported = service.import_file(image).document
    assert imported.extraction_status is ExtractionStatus.PENDING
    assert Path(imported.stored_path).is_file()

    with pytest.raises(ValueError, match="Timezone-aware"):
        repository.create(Document(
            None, "bad", "source", "stored", "TXT", "text/plain", 0, "a" * 64,
            datetime.now(), datetime.now(UTC), metadata_json=json.dumps({}),
        ))


def test_document_presenter_lists_and_imports_without_ui(document_services, tmp_path):
    repository, service = document_services
    presenter = DocumentPresenter(repository, service)
    source = tmp_path / "agenda.txt"; source.write_text("agenda", encoding="utf-8")
    presenter.import_file(source)
    assert presenter.list_documents() == ["agenda | TXT | EXTRACTED"]
